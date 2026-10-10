"""知识资产统一入口（入口②）：需求文档 / 技术文档 / 测试方案 / 历史缺陷 四类上传。

闭环设计（LightRAG / OpenWiki 取舍落地，见 design/loop-prototype/index.html）：
- 解析（parsers：txt/md/rst/csv/pdf/docx，历史缺陷另支持 xlsx）→ 敏感扫描（文本级，
  补齐 /api/kg/documents/upload 对 pdf/docx 不扫的缺口）→ 评估门卫 → 双写：
  ① rag_documents 业务 kind 行（即时可检索，生成上下文直接召回）
  ② LightRAG kg 异步管线（进图谱，workspace=repo:{id}——kind 不映射 workspace，
    否则跨 kind 连边与社区检测被打碎，LightRAG 原版取舍）
- 历史缺陷走「结构化表 + 摘要索引」：真相放 defects 表（可回归/流转/统计），
  KG 只做投影不逐行走 LLM 抽取（LightRAG ainsert_custom_kg 姿势的取舍版）。
- 需求文档走 requirements.ingest_one 四步管线（可测性评分 → G0 门禁）。
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import uuid

from fastapi import APIRouter, Request

from gateway.envelope import ApiError, ok

router = APIRouter()

ASSET_KINDS = {
    "req_doc": "需求文档",
    "tech_doc": "技术文档",
    "test_plan": "测试方案",
    "defect": "历史缺陷",
}

DOC_SUFFIXES = (".txt", ".md", ".markdown", ".rst", ".csv", ".pdf", ".docx")
DEFECT_SUFFIXES = (".csv", ".xlsx")

_SEVERITY_MAP = {"p0": "致命", "p1": "严重", "p2": "一般", "p3": "轻微", "blocker": "致命", "critical": "致命", "major": "严重", "minor": "一般", "trivial": "轻微"}
_DEFECT_COLUMNS = {
    "title": ("标题", "title", "摘要", "summary", "名称", "缺陷名称", "问题描述"),
    "detail": ("描述", "detail", "description", "内容", "复现步骤", "备注", "详情"),
    "severity": ("严重", "严重程度", "severity", "等级", "优先级", "priority", "级别"),
    "status": ("状态", "status"),
    "module": ("模块", "module", "组件"),
    "req_code": ("需求", "需求编号", "req_code", "req", "关联需求"),
}


def _require_admin(request: Request) -> None:
    from gateway.auth_routes import _require_admin as _guard

    _guard(request)


def _sensitive_guard(title: str, text: str) -> None:
    from services.shared.sensitive import scan_sensitive

    hits = scan_sensitive(f"{title}\n{text}")
    if hits:
        raise ApiError(422, f"文档命中敏感信息（{'、'.join(hits)}），已拒绝入库——请脱敏后重试", 422)


def _assess_guard(title: str, text: str) -> None:
    from gateway.main import _assess_knowledge_content

    score, reason = _assess_knowledge_content(title, text)
    if score < 0.5:
        raise ApiError(422, f"评估门卫：知识分 {score:.2f} < 0.5，拒绝入库——{reason}", 422)


def _submit_kg(workspace: str, title: str, text: str, repo_id: int) -> str:
    """投递 LightRAG kg 管线（异步：分块→抽取→合并→图谱），失败不阻塞资产入库。"""
    try:
        from services.shared.doc_pipeline import submit_document

        res = submit_document(workspace, title, text, repo_id=repo_id)
        return str(res.get("doc_key", ""))
    except Exception:  # noqa: BLE001
        return ""


def _ingest_doc_asset(kind: str, title: str, text: str, repo_id: int, filename: str) -> dict:
    """需求文档/技术文档/测试方案共用：门卫 → rag kind 行 → kg 管线。"""
    from gateway.requirements import ingest_one

    if kind == "req_doc":
        # 需求走四步管线（可测性评分即 OpenWiki 评估门卫的需求版，不再重复 LLM 评估）
        res = ingest_one(title, text, repo_id=repo_id, source="upload")
        _submit_kg(f"repo:{repo_id}" if repo_id else "default", f"需求 {res['code']} {title}", text, repo_id)
        return {"req_code": res["code"], "status": res["status"], "testability": res["testability"]}

    _assess_guard(title, text)
    digest = hashlib.sha1(f"{repo_id}:{title}".encode()).hexdigest()[:16]
    if kind == "tech_doc":
        doc_key, rag_kind = "userdoc:" + digest, "user_doc"
    else:  # test_plan
        doc_key, rag_kind = "plan:upload:" + digest, "plan"
    _index_asset(doc_key, rag_kind, title, text, repo_id, kind, filename)
    kg_key = _submit_kg(f"repo:{repo_id}" if repo_id else "default", title, text, repo_id)
    return {"doc_key": doc_key, "kg_doc_key": kg_key, "indexed": True}


def _index_asset(doc_key: str, rag_kind: str, title: str, text: str, repo_id: int, asset_kind: str, filename: str) -> None:
    from services.shared.docstatus import mark
    from services.shared.rag import ensure_rag_documents_table, index_document

    ensure_rag_documents_table()
    index_document(
        doc_key,
        rag_kind,
        title,
        text,
        repo_id=repo_id,
        meta={"source": "asset-upload", "asset_kind": asset_kind, "filename": filename},
    )
    mark(repo_id, rag_kind, doc_key)


@router.post("/api/knowledge/assets/upload")
async def upload_knowledge_asset(request: Request):
    """统一资产上传（multipart）：file + kind(req_doc|tech_doc|test_plan|defect) + repo_id。"""
    _require_admin(request)
    form = await request.form()
    upload = form.get("file")
    if upload is None or not hasattr(upload, "read"):
        raise ApiError(400, "multipart 字段 file 必填", 400)
    filename = str(getattr(upload, "filename", "doc.txt"))
    kind = str(form.get("kind") or "").strip()
    if kind not in ASSET_KINDS:
        raise ApiError(400, f"kind 必须是 {'/'.join(ASSET_KINDS)}", 400)
    repo_id = int(str(form.get("repo_id") or "0") or "0")
    title_override = str(form.get("title") or "").strip()
    data = await upload.read()
    if not data:
        raise ApiError(400, "上传文件为空", 400)

    suffix = ("." + filename.rsplit(".", 1)[-1].lower()) if "." in filename else ""
    if kind == "defect":
        if suffix not in DEFECT_SUFFIXES:
            raise ApiError(422, f"历史缺陷仅支持 {'/'.join(DEFECT_SUFFIXES)}（首行为表头）", 422)
        rows = _parse_defect_rows(filename, data)
        imported, codes = _import_defects(rows, repo_id, filename)
        return ok({"kind": kind, "imported": imported, "defect_codes": codes[:50], "filename": filename})

    if suffix not in DOC_SUFFIXES:
        raise ApiError(422, f"不支持的文件类型 {suffix or '(无后缀)'}（支持 txt/md/rst/csv/pdf/docx）", 422)

    from services.shared.parsers import parse_file

    try:
        parsed = parse_file(filename, data)
    except RuntimeError as exc:
        raise ApiError(422, str(exc), 422) from exc
    title = title_override or parsed.title or filename.rsplit(".", 1)[0]
    text = (parsed.text or "").strip()
    if len(text) < 30:
        raise ApiError(422, "解析出的文本不足 30 字（扫描件/图片型 PDF 请先转文字），无入库价值", 422)

    # 敏感门卫对解析后全文生效（含 pdf/docx，补齐 kg 管线只扫文本类的缺口）
    _sensitive_guard(title, text)
    res = _ingest_doc_asset(kind, title, text, repo_id, filename)
    return ok({"kind": kind, "title": title, "chars": len(text), "filename": filename, **res})


@router.get("/api/knowledge/assets/summary")
def assets_summary(repo_id: int = 0):
    """知识闭环总览：各资产 kind 计数 + 业务侧消费/回流计数（资产中心页顶部卡片）。"""
    from sqlalchemy import func
    from sqlalchemy import text as sql_text

    from gateway.main import get_session
    from services.shared.models import Cases, Defects, Functions, Requirements, WikiPages
    from services.shared.rag import ensure_rag_documents_table

    ensure_rag_documents_table()
    with get_session() as sess:
        q = sess.execute(
            sql_text("SELECT kind, COUNT(*) FROM rag_documents WHERE kind IN ('req','plan','user_doc','defect','lesson','wiki','case') GROUP BY kind"),
            {},
        )
        kinds = {k: int(n) for k, n in q.all()}
        rid = repo_id or None

        def _count(model, *conds):  # type: ignore[no-untyped-def]
            query = sess.query(func.count(model.id))
            if rid is not None and hasattr(model, "repo_id"):
                query = query.filter(model.repo_id == rid)
            for c in conds:
                query = query.filter(c)
            return int(query.scalar() or 0)

        summary = {
            "kinds": kinds,
            "requirements": _count(Requirements),
            "cases": _count(Cases),
            "wiki_pages": _count(WikiPages),
            "functions": _count(Functions),
            "defects_total": _count(Defects),
            "defects_open": _count(Defects, Defects.status != "已关闭"),
        }
    return ok(summary)


# ---------------- 历史缺陷结构化导入 ----------------


def _parse_defect_rows(filename: str, data: bytes) -> list[dict]:
    """csv/xlsx → 规范化行字典（列名按 _DEFECT_COLUMNS 模糊映射）。"""
    rows: list[dict] = []
    if filename.lower().endswith(".xlsx"):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ApiError(422, "xlsx 解析需要 openpyxl：uv add openpyxl", 422) from exc
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        raw = [[("" if c is None else str(c).strip()) for c in row] for row in ws.iter_rows(values_only=True)]  # type: ignore[union-attr]
        wb.close()
    else:
        text = data.decode("utf-8-sig", errors="replace")
        raw = [[c.strip() for c in r] for r in csv.reader(io.StringIO(text)) if any(c.strip() for c in r)]
    if not raw:
        return []
    header = raw[0]

    def col_of(field: str) -> int:
        for i, h in enumerate(header):
            if h.strip().lower() in [c.lower() for c in _DEFECT_COLUMNS[field]]:
                return i
        return -1

    idx = {f: col_of(f) for f in _DEFECT_COLUMNS}
    if idx["title"] < 0:
        raise ApiError(422, f"缺少标题列（表头需包含 {'/'.join(_DEFECT_COLUMNS['title'][:3])}…）", 422)
    for line in raw[1:]:
        row = {}
        for f, i in idx.items():
            row[f] = line[i].strip() if 0 <= i < len(line) else ""
        rows.append(row)
    return rows


def _import_defects(rows: list[dict], repo_id: int, filename: str) -> tuple[int, list[str]]:
    """规范化行 → defects 表（真相源）+ defect kind 摘要索引（生成 bugs 路立即召回）。"""
    from gateway.main import get_session
    from services.shared.models import Defects
    from services.shared.rag import ensure_rag_documents_table, index_document
    from services.shared.trace import emit

    ensure_rag_documents_table()
    imported, codes, skipped = 0, [], 0
    with get_session() as sess:
        for row in rows:
            title = (row.get("title") or "").strip()
            if not title:
                skipped += 1
                continue
            severity = _norm_severity(row.get("severity", ""))
            status = (row.get("status") or "新建").strip() or "新建"
            req_code = (row.get("req_code") or "").strip()
            if req_code and not re.match(r"^REQ-\d+$", req_code, re.IGNORECASE):
                req_code = ""
            module = (row.get("module") or "").strip()
            detail_txt = (row.get("detail") or "").strip()
            defect = Defects(
                code=f"BUG-{uuid.uuid4().hex[:6].upper()}",
                title=title[:500],
                origin_run="",
                case_codes="[]",
                req_code=req_code,
                severity=severity,
                status=status,
                assignee="未指派",
                trace_id="",
                detail=json.dumps({"source": "import", "filename": filename, "module": module, "description": detail_txt[:4000]}, ensure_ascii=False),
            )
            sess.add(defect)
            sess.flush()
            codes.append(defect.code)
            index_document(
                f"defect:{defect.code}",
                "defect",
                f"缺陷 {defect.code} {defect.title}",
                f"缺陷 {defect.code}（{severity}，状态：{status}）\n来源：历史缺陷导入（{filename}）\n关联需求：{req_code or '-'}\n模块：{module or '-'}\n描述：{detail_txt[:2000]}",
                repo_id=repo_id,
                meta={"source": "defect-import", "defect_code": defect.code, "status": status, "severity": severity},
            )
            imported += 1
        sess.commit()
    emit("缺陷", "asset-import", f"历史缺陷导入 {imported} 条（跳过空标题 {skipped}，来源 {filename}）")
    return imported, codes


def _norm_severity(raw: str) -> str:
    s = (raw or "").strip().lower()
    for key, val in _SEVERITY_MAP.items():
        if key in s:
            return val
    if s in ("致命", "严重", "一般", "轻微"):
        return s
    return "一般"
